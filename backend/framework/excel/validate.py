#!/usr/bin/env python3
"""验证层：重算只证明公式"能算"，不证明它"算对"。门禁分三级：

  L1 结构性：零公式错误（#REF!/#DIV/0!/…）、无小写未解析公式残留
             （LibreOffice 没能解析的函数会被小写写回——#NAME? 旁最快的判断依据）
  L2 契约性：页签名 / 表头 / 哪些单元格必须是公式而非硬编码值
  L3 数值性：抽样断言——指定单元格期望值（带容差）/ >0 / 非空

契约（contract dict）示例：
{
  "required_sheets": ["Summary", "Data"],
  "required_headers": {"Data": ["Date", "Region", "Revenue"]},
  "must_be_formula": ["Summary!B10", "Data!E2:E100"],
  "expect_values": [{"cell": "Summary!B10", "equals": 152340.5, "tol": 0.01},
                    {"cell": "Summary!B11", "gt": 0}],
  "max_formula_errors": 0
}
须跑在【影子重算后】的副本上——它才有真实缓存值（见 verify.py 链路）。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")

# 契约的全部合法键（max_changed 由 verify 层消费，也算合法）。未知键不再静默忽略——
# 小模型手滑键名（expected_values / require_sheets …）会让断言悄悄变成全过。
KNOWN_CONTRACT_KEYS = {"required_sheets", "required_headers", "must_be_formula",
                       "expect_values", "max_formula_errors", "max_changed"}


def _split_ref(ref: str) -> tuple:
    if "!" not in ref:
        raise ValueError(f"reference must be Sheet!A1 form: {ref}")
    sheet, rng = ref.split("!", 1)
    return sheet.strip("'"), rng


def _iter_cells(wb, ref: str):
    sheet, rng = _split_ref(ref)
    ws = wb[sheet]
    min_c, min_r, max_c, max_r = range_boundaries(rng)
    for r in range(min_r, (max_r or min_r) + 1):
        for c in range(min_c, (max_c or min_c) + 1):
            yield ws.cell(r, c)


def validate_workbook(path, contract: dict | None = None) -> dict:
    """L1/L2/L3 验证。返回 {"passed": bool, "checks_run": N, "failures": [...]}。"""
    p = Path(path)
    if not p.exists():
        return {"passed": False, "checks_run": 0,
                "failures": [{"level": "L0", "check": "file_exists", "error": str(p)}]}

    contract = contract or {}
    keep_vba = p.suffix.lower() == ".xlsm"
    wb_f = load_workbook(p, data_only=False, keep_vba=keep_vba)
    wb_v = load_workbook(p, data_only=True, keep_vba=keep_vba)

    failures: list = []
    checks = 0

    # ---------------- L1 结构性
    err_cells: list = []
    for name in wb_v.sheetnames:
        for row in wb_v[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERRORS:
                    err_cells.append(f"{name}!{cell.coordinate}={cell.value}")
    checks += 1
    max_err = contract.get("max_formula_errors", 0)
    if len(err_cells) > max_err:
        failures.append({"level": "L1", "check": "formula_errors",
                         "found": err_cells[:50], "total": len(err_cells)})

    # LibreOffice 无法解析的公式会被小写写回——快速信号
    lowered = []
    for name in wb_f.sheetnames:
        for row in wb_f[name].iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("=") and v[1:2].islower():
                    lowered.append(f"{name}!{cell.coordinate}")
    checks += 1
    if lowered:
        failures.append({"level": "L1", "check": "unparsed_lowercased_formula",
                         "found": lowered[:20], "total": len(lowered),
                         "hint": "LibreOffice 没能解析这些函数——多半是动态数组类函数"
                                 "或缺 _xlfn. 前缀，换等价的经典函数"})

    # ---------------- L2 契约性
    for sheet in contract.get("required_sheets", []):
        checks += 1
        if sheet not in wb_f.sheetnames:
            failures.append({"level": "L2", "check": "required_sheet", "missing": sheet,
                             "actual": wb_f.sheetnames})

    for sheet, headers in (contract.get("required_headers") or {}).items():
        checks += 1
        if sheet not in wb_f.sheetnames:
            failures.append({"level": "L2", "check": "required_headers",
                             "missing_sheet": sheet})
            continue
        ws = wb_f[sheet]
        actual = [ws.cell(1, c).value for c in range(1, len(headers) + 1)]
        if actual != headers:
            failures.append({"level": "L2", "check": "required_headers", "sheet": sheet,
                             "expected": headers, "actual": actual})

    for ref in contract.get("must_be_formula", []):
        checks += 1
        bad = []
        try:
            for cell in _iter_cells(wb_f, ref):
                v = cell.value
                if v in (None, ""):
                    continue
                if not (isinstance(v, str) and v.startswith("=")):
                    bad.append(f"{cell.coordinate}={v!r}")
        except (KeyError, ValueError) as exc:
            failures.append({"level": "L2", "check": "must_be_formula", "ref": ref,
                             "error": str(exc)})
            continue
        if bad:
            failures.append({"level": "L2", "check": "must_be_formula", "ref": ref,
                             "hardcoded": bad[:20], "total": len(bad),
                             "hint": "这些单元格被写死成了值，源数据变了不会重算"})

    # ---------------- L3 数值性
    for a in contract.get("expect_values", []):
        checks += 1
        ref = a.get("cell", "")
        try:
            cell = next(_iter_cells(wb_v, ref))
        except (KeyError, ValueError, StopIteration) as exc:
            failures.append({"level": "L3", "check": "expect_values", "cell": ref,
                             "error": str(exc)})
            continue
        got = cell.value
        ok, why = True, ""
        if "equals" in a:
            exp, tol = a["equals"], a.get("tol", 0)
            if isinstance(exp, (int, float)) and isinstance(got, (int, float)):
                ok = abs(got - exp) <= tol
            else:
                ok = got == exp
            why = f"expected {exp} (±{tol}), got {got!r}"
        if ok and "gt" in a:
            ok = isinstance(got, (int, float)) and got > a["gt"]
            why = f"expected > {a['gt']}, got {got!r}"
        if ok and "not_null" in a:
            ok = got not in (None, "")
            why = f"expected non-empty, got {got!r}"
        if not ok:
            failures.append({"level": "L3", "check": "expect_values", "cell": ref,
                             "why": why})

    result = {"file": str(p), "passed": not failures, "checks_run": checks,
              "failures": failures}
    unknown = [k for k in contract if k not in KNOWN_CONTRACT_KEYS]
    if unknown:
        result["unknown_contract_keys"] = unknown
        result["hint"] = (f"contract keys {unknown} were IGNORED — valid keys: "
                          f"{sorted(KNOWN_CONTRACT_KEYS)}")
    return result


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--contract")
    args = ap.parse_args()

    contract = json.loads(Path(args.contract).read_text()) if args.contract else None
    result = validate_workbook(args.path, contract)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
