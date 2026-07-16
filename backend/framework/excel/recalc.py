#!/usr/bin/env python3
"""重算层：openpyxl 写出的公式只有字符串、没有缓存值，必须过一遍 LibreOffice
headless 把整本 workbook 算一遍、就地重写文件，再扫描公式错误。

实现方式（沿用上游）：往一次性 LibreOffice profile 写 OOXMLRecalcMode=0（载入时
总是重算），再 --convert-to 出到临时目录、覆盖回原路径。不用 Basic 宏——
vnd.sun.star.script 那条路在无头环境里容易挂死等超时。

macOS 适配：
 · soffice 查找顺序 = PATH → /Applications/LibreOffice.app（brew cask 会把
   wrapper 链进 /opt/homebrew/bin，但 app 进程的 PATH 未必包含它）；
 · SAL_USE_VCLPLUGIN=svp 是 Linux headless 后端设置，darwin 上不设。

外部链接红线：`='[1]Sheet'!$B$2` 的缓存值是其数据的唯一载体，openpyxl 保存已
剥掉它，LibreOffice 会去解析磁盘上不存在的文件、失败写 #NAME? 并删光链接。
检测到外链一律拒绝（error_kind=external_links），除非 force=True。
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")
MAX_LOC = 100

REGISTRY = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
           xmlns:xs="http://www.w3.org/2001/XMLSchema">
  <item oor:path="/org.openoffice.Office.Calc/Formula/Load">
    <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Calc/Formula/Load">
    <prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop>
  </item>
</oor:items>
"""

FILTERS = {
    ".xlsx": "xlsx:Calc MS Excel 2007 XML",
    ".xlsm": "xlsm:Calc MS Excel 2007 VBA XML",
}

_SOFFICE_FALLBACKS = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
)


def find_soffice() -> Optional[str]:
    """定位**可用的** soffice；找不到返回 None（包壳工具据此决定 excel_research 是否注册）。

    实测事故（D1 验收）：brew cask 卸载 LibreOffice 后 `/opt/homebrew/bin/soffice`
    残链仍在，`which` 照样命中——工具带病注册，运行到 soffice 一步才炸 exit 127。
    探活必须真执行一次（--version），残链/坏装当场出局（扩充纪律：探活不是摆设）。
    """
    candidates = [p for name in ("soffice", "libreoffice") if (p := shutil.which(name))]
    candidates += [p for p in _SOFFICE_FALLBACKS if Path(p).is_file()]
    for p in candidates:
        try:
            r = subprocess.run([p, "--version"], capture_output=True, timeout=20)
            if r.returncode == 0:
                return p
            print(f"[excel] soffice 候选 {p} 探活失败（exit {r.returncode}）——跳过。"
                  "若为 brew 残链，可 `brew reinstall --cask libreoffice` 修复。")
        except Exception as e:
            print(f"[excel] soffice 候选 {p} 探活异常（{type(e).__name__}: {e}）——跳过。")
    return None


def _soffice_env(home: str) -> dict:
    env = {**os.environ, "HOME": home}
    if sys.platform == "linux":  # svp 是 Linux 的 headless VCL 后端；darwin 无此插件
        env["SAL_USE_VCLPLUGIN"] = "svp"
    return env


def _make_profile(root: Path) -> Path:
    prof = root / "prof"
    (prof / "user").mkdir(parents=True, exist_ok=True)
    (prof / "user" / "registrymodifications.xcu").write_text(REGISTRY, encoding="utf-8")
    return prof


def external_links(path) -> list:
    """返回 workbook 外部链接目标列表（空 = 安全）。"""
    wb = load_workbook(path, data_only=False)
    out = []
    for el in getattr(wb, "_external_links", []) or []:
        tgt = getattr(getattr(el, "file_link", None), "Target", None)
        if tgt:
            out.append(str(tgt))
    return out


def _run_recalc(path: Path, timeout: int) -> None:
    ext = path.suffix.lower()
    if ext not in FILTERS:
        raise RuntimeError(f"unsupported extension: {ext}")
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice not found (need `soffice` on PATH or in /Applications)")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        prof = _make_profile(root)
        outdir = root / "out"
        outdir.mkdir()
        cmd = [
            soffice, "--headless", "--norestore", "--invisible", "--nolockcheck",
            f"-env:UserInstallation=file://{prof}",
            "--convert-to", FILTERS[ext],
            "--outdir", str(outdir), str(path.resolve()),
        ]
        proc = subprocess.run(cmd, capture_output=True, timeout=timeout,
                              env=_soffice_env(tmp))
        produced = outdir / path.name
        if proc.returncode != 0 or not produced.exists():
            err = (proc.stderr or proc.stdout).decode(errors="replace")[:400]
            raise RuntimeError(f"soffice exit {proc.returncode}: {err}")
        shutil.move(str(produced), str(path))  # 就地覆盖


def _scan(path: Path) -> dict:
    keep_vba = path.suffix.lower() == ".xlsm"
    wb_f = load_workbook(path, data_only=False, keep_vba=keep_vba)
    wb_v = load_workbook(path, data_only=True, keep_vba=keep_vba)
    total_formulas = 0
    summary: dict = {}
    for name in wb_f.sheetnames:
        for row in wb_f[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total_formulas += 1
        for row in wb_v[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERRORS:
                    e = summary.setdefault(cell.value, {"count": 0, "locations": []})
                    e["count"] += 1
                    if len(e["locations"]) < MAX_LOC:
                        e["locations"].append(f"{name}!{cell.coordinate}")
    for e in summary.values():
        if e["count"] > len(e["locations"]):
            e["locations_truncated"] = e["count"] - len(e["locations"])
    total_errors = sum(e["count"] for e in summary.values())
    out = {
        "status": "errors_found" if total_errors else "success",
        "total_formulas": total_formulas,
        "total_errors": total_errors,
    }
    if summary:
        out["error_summary"] = summary
    return out


def recalc_workbook(path, *, timeout: int = 120, force: bool = False) -> dict:
    """就地重算并扫描公式错误。

    返回 dict：
      成功  {"status": "success"|"errors_found", "total_formulas", "total_errors", ...}
             ⚠ errors_found 也算"重算成功"——门禁判定交给 validate 层；
      失败  {"error": 给模型看的原因, "error_kind": external_links|timeout|recalc_failed}
             文件未被改动（外链拒绝在动文件之前）。
    """
    p = Path(path)
    if not p.exists():
        return {"error": f"file not found: {p}", "error_kind": "recalc_failed"}
    try:
        links = external_links(p)
        if links and not force:
            return {
                "error": ("workbook 含外部链接。openpyxl 保存时会剥掉这些单元格的缓存值，"
                          "LibreOffice 随后去解析磁盘上并不存在的那个文件，失败后写入 "
                          "#NAME? 并删掉所有链接。请先把这些单元格的值从原件抄出来"
                          "（对原件用 data_only=True 读），或 force 接受丢失。"),
                "error_kind": "external_links",
                "external_links": links,
            }
        _run_recalc(p, timeout)
        return _scan(p)
    except subprocess.TimeoutExpired:
        return {"error": f"recalc timed out after {timeout}s", "error_kind": "timeout"}
    except Exception as exc:  # noqa: BLE001
        return {"error": f"{type(exc).__name__}: {exc}", "error_kind": "recalc_failed"}


def warm() -> dict:
    """预热：把 LibreOffice 首次启动建 profile 那几秒提前烧掉（部署脚本可选跑一次）。"""
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        f = root / "warm.xlsx"
        wbk = Workbook()
        wbk.active["A1"] = 1
        wbk.active["A2"] = "=A1+1"
        wbk.save(f)
        result = recalc_workbook(f, timeout=240)
    return {"status": "warmed", "probe": result.get("status", result.get("error"))}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("path", nargs="?")
    ap.add_argument("--timeout", type=int, default=120)
    ap.add_argument("--force", action="store_true",
                    help="即使存在外部链接也强行重算（接受链接与缓存值丢失）")
    ap.add_argument("--warm", action="store_true", help="预热 LibreOffice profile")
    args = ap.parse_args()

    if args.warm:
        print(json.dumps(warm()))
        return 0
    if not args.path:
        print(json.dumps({"error": "usage: recalc.py <file.xlsx> [--timeout N] [--force]"}))
        return 2

    result = recalc_workbook(args.path, timeout=args.timeout, force=args.force)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if "error" not in result:
        return 0
    return {"external_links": 3, "timeout": 4}.get(result.get("error_kind"), 5)


if __name__ == "__main__":
    sys.exit(main())
