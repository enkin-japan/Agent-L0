#!/usr/bin/env python3
"""审计层（blast radius）：对比改动前后两个 workbook，输出单元格级变更清单。

两个用途：
 1. 门禁自检：只让改 D 列却动了 3000 格（典型于 pandas 读入再 to_excel 写回，
    格式和公式全被抹掉）——该路线最高频的静默破坏，必须拦住；
 2. 交付摘要："我改了这些"，比让人肉眼比对两个文件靠谱得多。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def _cell_sig(cell) -> tuple:
    """值/公式 + 关键样式。样式全量比对太吵，只取会影响可读性的几项。"""
    f = cell.font
    return (
        cell.value,
        cell.number_format,
        bool(f.bold),
        str(getattr(f.color, "rgb", None)) if f.color else None,
        str(getattr(cell.fill.fgColor, "rgb", None)) if cell.fill and cell.fill.fgColor else None,
    )


def diff_workbooks(a, b, *, cap: int = 100) -> dict:
    """逐格 diff（a=before, b=after）。changes 最多列 cap 条，counts 是全量计数。"""
    wa = load_workbook(a, data_only=False)
    wb = load_workbook(b, data_only=False)

    out: dict = {
        "before": str(a),
        "after": str(b),
        "sheets_added": [s for s in wb.sheetnames if s not in wa.sheetnames],
        "sheets_removed": [s for s in wa.sheetnames if s not in wb.sheetnames],
        "changes": [],
        "counts": {"value": 0, "formula": 0, "format": 0, "added": 0, "removed": 0},
    }

    for name in wb.sheetnames:
        if name not in wa.sheetnames:
            continue
        sa, sb = wa[name], wb[name]
        rows = max(sa.max_row, sb.max_row)
        cols = max(sa.max_column, sb.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                ca, cb = sa.cell(r, c), sb.cell(r, c)
                siga, sigb = _cell_sig(ca), _cell_sig(cb)
                if siga == sigb:
                    continue
                va, vb = siga[0], sigb[0]
                if va in (None, "") and vb not in (None, ""):
                    kind = "added"
                elif vb in (None, "") and va not in (None, ""):
                    kind = "removed"
                elif isinstance(vb, str) and vb.startswith("="):
                    kind = "formula"
                elif va != vb:
                    kind = "value"
                else:
                    kind = "format"
                out["counts"][kind] += 1
                if len(out["changes"]) < cap:
                    out["changes"].append(
                        {"cell": f"{name}!{cb.coordinate}", "kind": kind,
                         "before": str(va)[:60], "after": str(vb)[:60]})

    out["total_changed"] = sum(out["counts"].values())
    out["changes_truncated"] = max(0, out["total_changed"] - len(out["changes"]))
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("before")
    ap.add_argument("after")
    ap.add_argument("--cap", type=int, default=100, help="最多列出多少条变更")
    ap.add_argument("--max-changed", type=int, default=None,
                    help="改动单元格数超过此值即判定失败 (exit 1)")
    args = ap.parse_args()

    result = diff_workbooks(Path(args.before), Path(args.after), cap=args.cap)
    over = args.max_changed is not None and result["total_changed"] > args.max_changed
    if over:
        result["blast_radius_exceeded"] = args.max_changed
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 1 if over else 0


if __name__ == "__main__":
    sys.exit(main())
