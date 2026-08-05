#!/usr/bin/env python3
"""Excel 底盘基准：OfficeCLI 公式引擎 vs LibreOffice 影子重算（三波之 3，基准先行不动门禁）。

问的是一个问题：officecli 的自研公式引擎，对"沙箱 openpyxl 写出的裸公式文件"
（无缓存值——生产 verify 链的真实输入形态）算得对不对、算得多快，够不够格进
verify 底盘。判分基准是**手写真值**（每题 expected 在代码里手算写死，不经任何
引擎——两引擎都只是被考者，防裁判下场）。

三个测量面：
  ① 公式保真：~35 题覆盖 excel 沙箱实产形态（聚合/条件/查找/文本/日期/错误
     传播/链式依赖/跨表/OFFSET/INDIRECT），双引擎逐题对真值。
  ② 延迟：LO 重算全程 vs officecli 单次读取全程（同一文件各自拷贝）。
  ③ 视觉回读可用性：LO→PDF→pdftoppm 三级链 vs officecli view screenshot。

用法：backend/.venv/bin/python backend/benchmarks/officecli_eval/run_officecli_eval.py
结果落 results/（gitignore），结论人工写 BASELINE.md。
"""
import json
import shutil
import subprocess
import sys
import time
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BACKEND))

from framework.excel.recalc import find_soffice, recalc_workbook          # noqa: E402
from framework.excel.render import find_pdftoppm, render_workbook         # noqa: E402
from framework.office.officecli import find_officecli, run_cli            # noqa: E402

RESULTS = Path(__file__).parent / "results"

# ---- 题库：expected 全部手算写死（fit：为何贴合本项目见各题注释段） ----
# 形态覆盖=excel 沙箱实产：聚合统计/条件族/查找族/文本组装/日期/错误传播/链式依赖。
ERROR = "<ERROR>"          # 期望任何错误值（#DIV/0! / #NAME? / Err:xxx）
CASES = [
    ("sum_range",     "=SUM(Data!C2:C6)",                                    1150.5),
    ("average",       "=AVERAGE(Data!C2:C6)",                                230.1),
    ("minmax_spread", "=MAX(Data!C2:C6)-MIN(Data!C2:C6)",                    300),
    ("count_nums",    "=COUNT(Data!C2:C6)",                                  5),
    ("counta",        "=COUNTA(Data!A2:A6)",                                 5),
    ("if_nested",     '=IF(SUM(Data!C2:C3)>250,"hi",IF(1>0,"mid","lo"))',    "hi"),
    ("sumif",         '=SUMIF(Data!B2:B6,"East",Data!C2:C6)',                550.5),
    ("countif",       '=COUNTIF(Data!B2:B6,"East")',                         3),
    ("sumifs",        '=SUMIFS(Data!C2:C6,Data!B2:B6,"East",Data!C2:C6,">120")', 450.5),
    ("vlookup",       '=VLOOKUP("Carol",Data!A2:C6,3,FALSE)',                300),
    ("index_match",   '=INDEX(Data!C2:C6,MATCH("Dave",Data!A2:A6,0))',       400),
    ("xlookup",       '=XLOOKUP("Bob",Data!A2:A6,Data!C2:C6)',               200),
    ("round_half",    "=ROUND(Data!C6,0)",                                   151),
    ("int_mod",       "=MOD(INT(Data!C4),7)",                                6),
    ("abs_power",     "=ABS(-3)+POWER(2,5)",                                 35),
    ("concat_amp",    '="R-"&Data!A2&"-"&COUNTIF(Data!B2:B6,"East")',        "R-Alice-3"),
    ("textjoin",      '=TEXTJOIN(",",TRUE,Data!A2:A3)',                      "Alice,Bob"),
    ("left_len",      "=LEFT(Data!A3,2)&LEN(Data!A4)",                       "Bo5"),
    ("upper_trim",    '=UPPER(TRIM("  ok  "))',                              "OK"),
    ("text_percent",  '=TEXT(0.256,"0.0%")',                                 "25.6%"),
    ("substitute",    '=SUBSTITUTE("a-b-c","-",".")',                        "a.b.c"),
    ("date_year",     "=YEAR(DATE(2026,8,5))",                               2026),
    ("eomonth_day",   "=DAY(EOMONTH(DATE(2026,2,10),0))",                    28),
    ("weekday_iso",   "=WEEKDAY(DATE(2026,8,5),2)",                          3),
    ("div_zero",      "=1/0",                                                ERROR),
    ("iferror_safe",  '=IFERROR(1/0,"safe")',                                "safe"),
    ("isna_catch",    '=IF(ISNA(NA()),"caught","no")',                       "caught"),
    ("name_error",    "=NOSUCHFN(1)",                                        ERROR),
    ("chain_local",   "=E29+1",                                              11),      # D29=5,E29==D29*2
    ("chain_cross",   "=SUM(Data!C2:C2)+E29",                                110),
    ("ratio_float",   "=Data!C3/SUM(Data!C2:C3)",                            200/300),
    ("sum_200",       "=SUM(G1:G200)",                                       20100),
    ("sumproduct",    "=SUMPRODUCT(Data!C2:C3,Data!C4:C5)",                  110000),
    ("offset_sum",    "=SUM(OFFSET(Data!C2,0,0,3,1))",                       600),
    ("indirect_ref",  '=INDIRECT("Data!C5")',                                400),
]


def build_workbook(path: Path):
    """openpyxl 写裸公式（无缓存值）——生产 verify 链的真实输入形态。"""
    from openpyxl import Workbook
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    rows = [("name", "region", "sales"), ("Alice", "East", 100), ("Bob", "West", 200),
            ("Carol", "East", 300), ("Dave", "South", 400), ("Erin", "East", 150.5)]
    for r in rows:
        data.append(r)
    f = wb.create_sheet("F")
    f["D29"], f["E29"] = 5, "=D29*2"
    for i in range(1, 201):
        f.cell(row=i, column=7, value=i)                    # G1:G200
    for idx, (cid, formula, _exp) in enumerate(CASES, 1):
        f.cell(row=idx, column=1, value=cid)
        f.cell(row=idx, column=2, value=formula)
    wb.save(path)


def _match(expected, actual) -> bool:
    if actual is None:
        return False
    s = str(actual)
    if expected is ERROR:
        return s.startswith("#") or s.startswith("Err") or s == "ERROR"
    if isinstance(expected, (int, float)):
        try:
            return abs(float(s) - float(expected)) <= max(1e-6, abs(expected) * 1e-9)
        except ValueError:
            return False
    return s == str(expected)


def read_lo(path: Path) -> dict:
    """LO 影子重算（生产 recalc 同函数）→ openpyxl data_only 读缓存值。"""
    from openpyxl import load_workbook
    t0 = time.time()
    r = recalc_workbook(path)
    dt = time.time() - t0
    # status=errors_found 是正常路径：题库故意含错误公式（div_zero/name_error），
    # 重算本身已完成，逐题对真值判分；status=fail 类才是重算失败。
    wb = load_workbook(path, data_only=True)
    f = wb["F"]
    return {"elapsed": dt, "status": r.get("status"),
            "values": {cid: f.cell(row=i, column=2).value
                       for i, (cid, _, _) in enumerate(CASES, 1)}}


def read_oc(path: Path) -> dict:
    """officecli 单次整表读取：get /F --depth 2（sheet→row→cell），取每格 computedValue。"""
    t0 = time.time()
    out = run_cli(["get", str(path), "/F", "--depth", "2"], timeout=120)
    dt = time.time() - t0
    cells = {}
    for row in out["data"]["results"][0].get("children", []):
        for node in row.get("children", []):
            addr = node.get("preview") or node.get("path", "").rsplit("/", 1)[-1]
            fmt = node.get("format") or {}
            val = fmt.get("computedValue", fmt.get("cachedValue", node.get("text")))
            if fmt.get("type") == "Error" and not str(val).startswith(("#", "Err")):
                val = f"#{val}"
            cells[addr] = val
    return {"elapsed": dt,
            "values": {cid: cells.get(f"B{i}")
                       for i, (cid, _, _) in enumerate(CASES, 1)}}


def visual_probe(path: Path, outdir: Path) -> dict:
    res = {}
    if find_soffice() and find_pdftoppm():
        t0 = time.time()
        try:
            pngs = render_workbook(path, outdir / "lo_png")
            res["lo"] = {"elapsed": time.time() - t0, "pages": len(pngs)}
        except Exception as e:
            res["lo"] = {"error": str(e)[:200]}
    else:
        res["lo"] = {"error": "soffice/pdftoppm missing"}
    t0 = time.time()
    p = subprocess.run([find_officecli(), "view", str(path), "screenshot",
                       "-o", str(outdir / "oc.png"), "--json"],
                      capture_output=True, text=True, timeout=120)
    res["oc"] = ({"elapsed": time.time() - t0}
                 if p.returncode == 0 else {"error": p.stderr[:200]})
    return res


def main() -> int:
    if find_officecli() is None:
        print("officecli not installed")
        return 1
    RESULTS.mkdir(exist_ok=True)
    src = RESULTS / "probe.xlsx"
    build_workbook(src)
    lo_file, oc_file = RESULTS / "lo.xlsx", RESULTS / "oc.xlsx"
    shutil.copy2(src, lo_file)
    shutil.copy2(src, oc_file)

    lo = read_lo(lo_file)
    oc = read_oc(oc_file)

    rows, agree = [], {"lo": 0, "oc": 0}
    for cid, formula, exp in CASES:
        lv, ov = lo["values"].get(cid), oc["values"].get(cid)
        lok, ook = _match(exp, lv), _match(exp, ov)
        agree["lo"] += lok
        agree["oc"] += ook
        rows.append({"id": cid, "formula": formula, "expected": exp if exp is not ERROR else "ERROR",
                     "lo": {"value": lv, "ok": lok}, "oc": {"value": ov, "ok": ook}})
        flag = "  " if (lok and ook) else "❌"
        print(f"{flag} {cid:14s} exp={str(exp)[:18]:18s} "
              f"LO={'✓' if lok else str(lv)[:20]:>20s} OC={'✓' if ook else str(ov)[:20]:>20s}")

    n = len(CASES)
    vis = visual_probe(src, RESULTS)
    summary = {"cases": n, "lo_pass": agree["lo"], "oc_pass": agree["oc"],
               "lo_recalc_s": round(lo["elapsed"], 2), "oc_read_s": round(oc["elapsed"], 2),
               "visual": vis}
    print(f"\nLO {agree['lo']}/{n} ({lo['elapsed']:.1f}s recalc) | "
          f"OC {agree['oc']}/{n} ({oc['elapsed']:.1f}s read) | visual={vis}")
    (RESULTS / "results.json").write_text(
        json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False, indent=1),
        encoding="utf-8")
    print(f"written: {RESULTS / 'results.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
